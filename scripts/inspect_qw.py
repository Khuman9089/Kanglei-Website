import openpyxl
import sys

filepath = r'C:\Users\MayNard\Desktop\Kuthi\qw.xlsm'

print("Loading workbook (data_only=True)...")
wb_data = openpyxl.load_workbook(filepath, read_only=True, data_only=True)

def inspect_sheet(wb, sheet_name, max_r=25, max_c=12):
    if sheet_name not in wb.sheetnames:
        print(f"Sheet {sheet_name} not found.")
        return
    ws = wb[sheet_name]
    print(f"\n--- Sheet: {sheet_name} (first {max_r} rows, {max_c} cols) ---")
    row_count = 0
    for row in ws.iter_rows(max_row=max_r, max_col=max_c, values_only=True):
        row_count += 1
        # only print non-empty rows
        if any(v is not None for v in row):
            cells = [str(v) if v is not None else "" for v in row]
            print(f"R{row_count:02d}: " + " | ".join(cells[:8]))

for s in ['Edit', 'Calculation', 'Paring Ka Yatpa', 'Meetei_Result', 'MM_Result_M', 'Sinpham', 'HOME']:
    inspect_sheet(wb_data, s, max_r=20, max_c=8)
